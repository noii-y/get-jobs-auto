"""投递结果与黑名单API"""
import csv
import io
from typing import Optional
from datetime import datetime
from urllib.parse import quote
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session
from app.database import get_db
from app.models import JobResult, Blacklist
from app.schemas import BlacklistCreate, BlacklistResponse
from app.services.blacklist import BlacklistService

router = APIRouter(prefix="/api", tags=["结果与黑名单"])


# ============ 投递结果 ============

@router.get("/results", summary="获取投递结果列表")
async def list_results(run_id: Optional[int] = None,
                       platform: Optional[str] = None,
                       status: Optional[str] = None,
                       limit: int = Query(50, le=500),
                       offset: int = 0,
                       db: Session = Depends(get_db)):
    """获取投递结果记录"""
    query = db.query(JobResult)
    if run_id:
        query = query.filter(JobResult.run_id == run_id)
    if platform:
        query = query.filter(JobResult.platform == platform)
    if status:
        query = query.filter(JobResult.status == status)

    total = query.count()
    results = query.order_by(JobResult.created_at.desc()).offset(offset).limit(limit).all()

    return {
        "total": total,
        "items": [
            {
                "id": r.id,
                "run_id": r.run_id,
                "platform": r.platform,
                "job_title": r.job_title,
                "company": r.company,
                "salary": r.salary,
                "city": r.city,
                "recruiter": r.recruiter,
                "job_url": r.job_url,
                "status": r.status,
                "message": r.message,
                "ai_greeting": r.ai_greeting,
                "reason_code": getattr(r, "reason_code", "") or "",
                "error_stage": getattr(r, "error_stage", "") or "",
                "error_url": getattr(r, "error_url", "") or "",
                "screenshot_path": getattr(r, "screenshot_path", "") or "",
                "selector_used": getattr(r, "selector_used", "") or "",
                "traceback_short": getattr(r, "traceback_short", "") or "",
                "created_at": r.created_at
            }
            for r in results
        ]
    }


@router.get("/results/stats", summary="获取投递统计")
async def get_stats(platform: Optional[str] = None,
                    days: int = Query(7, description="统计最近N天"),
                    db: Session = Depends(get_db)):
    """获取投递统计数据"""
    from datetime import datetime, timedelta
    since = datetime.now() - timedelta(days=days)

    query = db.query(JobResult).filter(JobResult.created_at >= since)
    if platform:
        query = query.filter(JobResult.platform == platform)

    results = query.all()
    total = len(results)
    success = sum(1 for r in results if r.status == "success")
    failed = sum(1 for r in results if r.status == "failed")
    skipped = sum(1 for r in results if r.status == "skipped")

    # 按平台统计
    platform_stats = {}
    for r in results:
        if r.platform not in platform_stats:
            platform_stats[r.platform] = {"total": 0, "success": 0, "failed": 0, "skipped": 0}
        platform_stats[r.platform]["total"] += 1
        platform_stats[r.platform][r.status if r.status in ("success", "failed", "skipped") else "failed"] += 1

    return {
        "period_days": days,
        "total": total,
        "success": success,
        "failed": failed,
        "skipped": skipped,
        "success_rate": round(success / total * 100, 1) if total > 0 else 0,
        "by_platform": platform_stats
    }


@router.get("/results/trend", summary="获取投递趋势数据")
async def get_trend(platform: Optional[str] = None,
                    days: int = Query(7, description="统计最近N天"),
                    db: Session = Depends(get_db)):
    """获取按天投递趋势和按小时分布数据"""
    from datetime import datetime, timedelta

    since = datetime.now() - timedelta(days=days)
    query = db.query(JobResult).filter(JobResult.created_at >= since)
    if platform:
        query = query.filter(JobResult.platform == platform)
    results = query.all()

    # 按天统计
    daily_data = {}
    for i in range(days):
        day = (datetime.now() - timedelta(days=days - 1 - i)).strftime("%Y-%m-%d")
        daily_data[day] = {"total": 0, "success": 0, "failed": 0, "skipped": 0}

    for r in results:
        if r.created_at:
            day = r.created_at.strftime("%Y-%m-%d")
            if day in daily_data:
                daily_data[day]["total"] += 1
                status = r.status if r.status in ("success", "failed", "skipped") else "failed"
                daily_data[day][status] += 1

    daily = [{"date": k, **v} for k, v in daily_data.items()]

    # 按小时统计（最佳投递时间）
    hourly_data = {h: {"total": 0, "success": 0} for h in range(24)}
    for r in results:
        if r.created_at:
            hour = r.created_at.hour
            hourly_data[hour]["total"] += 1
            if r.status == "success":
                hourly_data[hour]["success"] += 1

    hourly = [
        {
            "hour": h,
            "total": hourly_data[h]["total"],
            "success": hourly_data[h]["success"],
            "success_rate": round(
                hourly_data[h]["success"] / hourly_data[h]["total"] * 100, 1
            ) if hourly_data[h]["total"] > 0 else 0
        }
        for h in range(24)
    ]

    # 找出最佳投递时段（成功率最高的3个小时，且投递数>=3）
    valid_hours = [h for h in hourly if h["total"] >= 3]
    best_hours = sorted(valid_hours, key=lambda x: x["success_rate"], reverse=True)[:3]

    return {
        "period_days": days,
        "daily": daily,
        "hourly": hourly,
        "best_hours": best_hours,
    }


# ============ 黑名单 ============

@router.get("/blacklist", summary="获取黑名单列表")
async def list_blacklist(platform: Optional[str] = None,
                         limit: int = Query(100, le=500),
                         offset: int = 0,
                         db: Session = Depends(get_db)):
    """获取企业黑名单"""
    service = BlacklistService(db)
    items = service.list_all(platform=platform or "", offset=offset, limit=limit)
    return {
        "total": service.count(),
        "items": [
            {
                "id": item.id,
                "company_name": item.company_name,
                "platform": item.platform,
                "reason": item.reason,
                "source": item.source,
                "created_at": item.created_at
            }
            for item in items
        ]
    }


@router.post("/blacklist", summary="添加黑名单")
async def add_blacklist(item: BlacklistCreate, db: Session = Depends(get_db)):
    """添加企业到黑名单"""
    service = BlacklistService(db)
    success = service.add(
        company_name=item.company_name,
        platform=item.platform,
        reason=item.reason,
        source="manual"
    )
    if not success:
        raise HTTPException(status_code=400, detail="添加失败，企业已存在或名称无效")
    return {"status": "ok", "company_name": item.company_name}


@router.delete("/blacklist/{blacklist_id}", summary="删除黑名单")
async def delete_blacklist(blacklist_id: int, db: Session = Depends(get_db)):
    """从黑名单移除企业"""
    service = BlacklistService(db)
    success = service.remove(blacklist_id)
    if not success:
        raise HTTPException(status_code=404, detail="记录不存在")
    return {"status": "ok"}


# ============ 导出 ============

def _query_results(db, run_id=None, platform=None, status=None):
    """查询投递结果（通用）"""
    query = db.query(JobResult)
    if run_id:
        query = query.filter(JobResult.run_id == run_id)
    if platform:
        query = query.filter(JobResult.platform == platform)
    if status:
        query = query.filter(JobResult.status == status)
    return query.order_by(JobResult.created_at.desc()).all()


@router.get("/results/export/csv", summary="导出投递记录为CSV")
async def export_csv(run_id: Optional[int] = None,
                     platform: Optional[str] = None,
                     status: Optional[str] = None,
                     db: Session = Depends(get_db)):
    """导出投递记录为CSV文件"""
    results = _query_results(db, run_id, platform, status)

    # 生成CSV
    output = io.StringIO()
    # 写入BOM以便Excel正确识别中文
    output.write('\ufeff')
    writer = csv.writer(output)
    writer.writerow(["ID", "平台", "职位名称", "公司", "薪资", "城市", "招聘者",
                     "状态", "消息", "AI问候语", "职位链接", "投递时间"])

    status_map = {"success": "成功", "failed": "失败", "skipped": "跳过"}
    platform_map = {"boss": "Boss直聘", "job51": "51job"}

    for r in results:
        writer.writerow([
            r.id,
            platform_map.get(r.platform, r.platform),
            r.job_title,
            r.company,
            r.salary,
            r.city,
            r.recruiter or "",
            status_map.get(r.status, r.status),
            r.message or "",
            r.ai_greeting or "",
            r.job_url or "",
            r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
        ])

    csv_content = output.getvalue()
    filename = f"投递记录_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    encoded_filename = quote(filename)

    return Response(
        content=csv_content,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )


@router.get("/results/export/excel", summary="导出投递记录为Excel")
async def export_excel(run_id: Optional[int] = None,
                       platform: Optional[str] = None,
                       status: Optional[str] = None,
                       db: Session = Depends(get_db)):
    """导出投递记录为Excel文件"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="未安装openpyxl，请先运行 pip install openpyxl")

    results = _query_results(db, run_id, platform, status)

    wb = Workbook()
    ws = wb.active
    ws.title = "投递记录"

    # 表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    headers = ["ID", "平台", "职位名称", "公司", "薪资", "城市", "招聘者",
               "状态", "消息", "AI问候语", "职位链接", "投递时间"]
    ws.append(headers)

    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    status_map = {"success": "成功", "failed": "失败", "skipped": "跳过"}
    platform_map = {"boss": "Boss直聘", "job51": "51job"}

    for r in results:
        ws.append([
            r.id,
            platform_map.get(r.platform, r.platform),
            r.job_title,
            r.company,
            r.salary,
            r.city,
            r.recruiter or "",
            status_map.get(r.status, r.status),
            r.message or "",
            r.ai_greeting or "",
            r.job_url or "",
            r.created_at.strftime("%Y-%m-%d %H:%M:%S") if r.created_at else "",
        ])

    # 自动调整列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width

    # 冻结首行
    ws.freeze_panes = "A2"

    # 保存到内存
    output = io.BytesIO()
    wb.save(output)
    excel_content = output.getvalue()

    filename = f"投递记录_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    encoded_filename = quote(filename)

    return Response(
        content=excel_content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"}
    )
